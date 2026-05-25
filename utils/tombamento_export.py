from __future__ import annotations

from copy import copy
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def _copy_row_style(worksheet, source_row: int, target_row: int, columns: range) -> None:
    for column in columns:
        source_cell = worksheet.cell(row=source_row, column=column)
        target_cell = worksheet.cell(row=target_row, column=column)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.font:
            target_cell.font = copy(source_cell.font)


def build_tombamento_export(template_path: Path, rebalance_detail: pd.DataFrame, cliente: str) -> bytes:
    workbook = load_workbook(template_path)
    worksheet = workbook["Importação"] if "Importação" in workbook.sheetnames else workbook.active

    output_rows = (
        rebalance_detail[rebalance_detail["quantidade_meta"] > 0]
        .loc[:, ["ticker", "quantidade_meta"]]
        .sort_values("ticker")
        .reset_index(drop=True)
    )

    start_row = 3
    data_columns = range(2, 5)
    template_style_row = start_row

    for row in range(start_row, max(worksheet.max_row, start_row + len(output_rows) + 5) + 1):
        for column in data_columns:
            worksheet.cell(row=row, column=column).value = None

    for index, row in output_rows.iterrows():
        excel_row = start_row + index
        if excel_row != template_style_row:
            _copy_row_style(worksheet, template_style_row, excel_row, data_columns)
        worksheet.cell(row=excel_row, column=2).value = str(cliente).strip()
        worksheet.cell(row=excel_row, column=3).value = row["ticker"]
        worksheet.cell(row=excel_row, column=4).value = int(row["quantidade_meta"])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
