from __future__ import annotations

from copy import copy
from io import BytesIO
from pathlib import Path
import unicodedata

import pandas as pd
from openpyxl import load_workbook


def _column_width(df: pd.DataFrame, column_position: int, header: object, minimum: int = 12, maximum: int = 35) -> int:
    values = [len(str(header))]
    if df is not None and not df.empty and column_position < len(df.columns):
        series = df.iloc[:, column_position]
        for item in series.head(200).tolist():
            values.append(0 if pd.isna(item) else len(str(item)))
    return max(minimum, min(maximum, max(values) + 2))


def build_excel_export(
    orders: pd.DataFrame,
    detail: pd.DataFrame,
    costs: dict[str, float],
    origem: pd.DataFrame,
    destino: pd.DataFrame,
    prices: pd.DataFrame,
    fiscal_detail: pd.DataFrame | None = None,
    fiscal_summary: dict[str, float | str] | None = None,
    fiscal_notice: str | None = None,
    ordem_cio_bytes: bytes | None = None,
    tombamento_bytes: bytes | None = None,
    audit_summary: dict[str, object] | None = None,
    audit_checks: pd.DataFrame | None = None,
) -> bytes:
    output = BytesIO()
    costs_df = pd.DataFrame(
        [
            {"Métrica": "Total compras", "Valor": costs["total_compras"]},
            {"Métrica": "Total vendas", "Valor": costs["total_vendas"]},
            {"Métrica": "Movimentação total", "Valor": costs["movimentacao_total"]},
            {"Métrica": "Corretagem 0,5%", "Valor": costs["corretagem"]},
            {"Métrica": "Custos adicionais 0,2%", "Valor": costs["custos_adicionais"]},
            {"Métrica": "Custo total estimado 0,7%", "Valor": costs["custo_total"]},
        ]
    )

    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        sheets = {
            "Ordens": orders,
            "Rebalanceamento Detalhado": detail,
            "Custos": costs_df,
            "Carteira Origem": origem,
            "Carteira Destino": destino,
            "Preços": prices,
        }

        workbook = writer.book
        header_format = workbook.add_format({"bold": True, "bg_color": "#D9EAF7", "border": 1})
        money_format = workbook.add_format({"num_format": 'R$ #,##0.00'})
        percent_format = workbook.add_format({"num_format": "0.00%"})
        title_format = workbook.add_format({"bold": True, "font_size": 13, "bg_color": "#EAF4EA"})
        wrap_format = workbook.add_format({"text_wrap": True, "valign": "top"})

        for sheet_name, df in sheets.items():
            safe_name = sheet_name[:31]
            df.to_excel(writer, index=False, sheet_name=safe_name)
            worksheet = writer.sheets[safe_name]
            for col_num, value in enumerate(df.columns):
                worksheet.write(0, col_num, value, header_format)
                width = _column_width(df, col_num, value)
                worksheet.set_column(col_num, col_num, width)

            for col_num, value in enumerate(df.columns):
                lowered = str(value).lower()
                if "valor" in lowered or "preco" in lowered or "financeiro" in lowered or "total" in lowered or value == "Valor":
                    worksheet.set_column(col_num, col_num, 16, money_format)
                if "peso" in lowered:
                    worksheet.set_column(col_num, col_num, 14, percent_format)

        if fiscal_detail is not None and fiscal_summary is not None:
            sheet_name = "Estimativa Fiscal"
            fiscal_detail.to_excel(writer, index=False, sheet_name=sheet_name, startrow=1)
            worksheet = writer.sheets[sheet_name]
            worksheet.write(0, 0, "Resultado fiscal por ativo vendido", title_format)

            for col_num, value in enumerate(fiscal_detail.columns):
                worksheet.write(1, col_num, value, header_format)
                width = _column_width(fiscal_detail, col_num, value)
                worksheet.set_column(col_num, col_num, width)
                lowered = str(value).lower()
                if (
                    "valor" in lowered
                    or "preco" in lowered
                    or "custo" in lowered
                    or "lucro" in lowered
                    or "ir" in lowered
                    or "prejuizo" in lowered
                ):
                    worksheet.set_column(col_num, col_num, 16, money_format)

            summary_df = pd.DataFrame(
                [{"Métrica": key, "Valor": value} for key, value in fiscal_summary.items()]
            )
            start_row = len(fiscal_detail) + 5
            worksheet.write(start_row, 0, "Resumo fiscal", title_format)
            summary_df.to_excel(writer, index=False, sheet_name=sheet_name, startrow=start_row + 1)
            for col_num, value in enumerate(summary_df.columns):
                worksheet.write(start_row + 1, col_num, value, header_format)
            worksheet.set_column(0, 0, 34)
            worksheet.set_column(1, 1, 24, money_format)

            notice_row = start_row + len(summary_df) + 4
            worksheet.write(notice_row, 0, "Avisos fiscais", title_format)
            worksheet.write(notice_row + 1, 0, fiscal_notice or "", wrap_format)
            worksheet.set_column(0, max(len(fiscal_detail.columns) - 1, 1), 20)
            worksheet.set_row(notice_row + 1, 70)

        if audit_summary is not None and audit_checks is not None:
            sheet_name = "Auditoria"
            worksheet = workbook.add_worksheet(sheet_name)
            writer.sheets[sheet_name] = worksheet
            worksheet.write(0, 0, "Resumo da auditoria", title_format)
            summary_df = pd.DataFrame(
                [{"Item": key, "Valor": value} for key, value in audit_summary.items()]
            )
            summary_df.to_excel(writer, index=False, sheet_name=sheet_name, startrow=1)
            for col_num, value in enumerate(summary_df.columns):
                worksheet.write(1, col_num, value, header_format)
            worksheet.set_column(0, 0, 34)
            worksheet.set_column(1, 1, 28)

            checks_start_row = len(summary_df) + 5
            worksheet.write(checks_start_row, 0, "Checks detalhados", title_format)
            audit_checks.to_excel(writer, index=False, sheet_name=sheet_name, startrow=checks_start_row + 1)
            for col_num, value in enumerate(audit_checks.columns):
                worksheet.write(checks_start_row + 1, col_num, value, header_format)
                width = _column_width(audit_checks, col_num, value, minimum=14, maximum=45)
                worksheet.set_column(col_num, col_num, width)

    excel_bytes = output.getvalue()
    if ordem_cio_bytes or tombamento_bytes:
        workbook = load_workbook(BytesIO(excel_bytes))
        if ordem_cio_bytes:
            source_workbook = load_workbook(BytesIO(ordem_cio_bytes))
            _copy_sheet_to_workbook(source_workbook.active, workbook, "Ordem CIO")
        if tombamento_bytes:
            source_workbook = load_workbook(BytesIO(tombamento_bytes))
            source_sheet = source_workbook["Importação"] if "Importação" in source_workbook.sheetnames else source_workbook.active
            _copy_sheet_to_workbook(source_sheet, workbook, "Tombamento")
        final_output = BytesIO()
        workbook.save(final_output)
        return final_output.getvalue()

    return excel_bytes


def _unique_sheet_title(workbook, title: str) -> str:
    clean_title = title[:31]
    if clean_title not in workbook.sheetnames:
        return clean_title
    for index in range(2, 100):
        suffix = f" {index}"
        candidate = f"{clean_title[:31 - len(suffix)]}{suffix}"
        if candidate not in workbook.sheetnames:
            return candidate
    return clean_title[:28] + " 99"


def _copy_sheet_to_workbook(source_ws, target_wb, title: str):
    target_ws = target_wb.create_sheet(_unique_sheet_title(target_wb, title))
    target_ws.sheet_format = copy(source_ws.sheet_format)
    target_ws.sheet_properties = copy(source_ws.sheet_properties)
    target_ws.page_margins = copy(source_ws.page_margins)
    target_ws.page_setup = copy(source_ws.page_setup)
    target_ws.print_options = copy(source_ws.print_options)
    target_ws.freeze_panes = source_ws.freeze_panes

    if source_ws.auto_filter:
        target_ws.auto_filter.ref = source_ws.auto_filter.ref

    for row in source_ws.iter_rows():
        for source_cell in row:
            target_cell = target_ws.cell(row=source_cell.row, column=source_cell.column, value=source_cell.value)
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            if source_cell.alignment:
                target_cell.alignment = copy(source_cell.alignment)
            if source_cell.border:
                target_cell.border = copy(source_cell.border)
            if source_cell.fill:
                target_cell.fill = copy(source_cell.fill)
            if source_cell.font:
                target_cell.font = copy(source_cell.font)
            if source_cell.protection:
                target_cell.protection = copy(source_cell.protection)
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)

    for merged_range in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged_range))

    for key, dimension in source_ws.column_dimensions.items():
        target_ws.column_dimensions[key].width = dimension.width
        target_ws.column_dimensions[key].hidden = dimension.hidden
        target_ws.column_dimensions[key].outlineLevel = dimension.outlineLevel
        target_ws.column_dimensions[key].collapsed = dimension.collapsed

    for key, dimension in source_ws.row_dimensions.items():
        target_ws.row_dimensions[key].height = dimension.height
        target_ws.row_dimensions[key].hidden = dimension.hidden
        target_ws.row_dimensions[key].outlineLevel = dimension.outlineLevel
        target_ws.row_dimensions[key].collapsed = dimension.collapsed

    return target_ws


def _normalize_header(value: object) -> str:
    text = "" if value is None else str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return "".join(char for char in text if char.isalnum())


def _find_ordem_cio_headers(ws) -> tuple[int, dict[str, int]]:
    expected = {
        "estrategia": "Estratégia",
        "cliente": "Cliente",
        "ativo": "Ativo",
        "cv": "C/V",
        "preco": "Preço",
        "qtdtotal": "Qtd. Total",
    }
    for row in range(1, ws.max_row + 1):
        found: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            key = _normalize_header(ws.cell(row=row, column=col).value)
            if key in expected:
                found[key] = col
        if set(expected).issubset(found):
            return row, found
    missing = ", ".join(expected.values())
    raise ValueError(f"Cabeçalhos obrigatórios não encontrados no modelo: {missing}")


def _get_order_value(row: pd.Series, aliases: list[str]) -> object:
    normalized = {_normalize_header(column): column for column in row.index}
    for alias in aliases:
        column = normalized.get(_normalize_header(alias))
        if column is not None:
            return row[column]
    return None


def _copy_row_style(ws, source_row: int, target_row: int, columns: list[int]) -> None:
    if source_row == target_row:
        return
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in columns:
        source = ws.cell(row=source_row, column=col)
        target = ws.cell(row=target_row, column=col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.border:
            target.border = copy(source.border)
        if source.fill:
            target.fill = copy(source.fill)
        if source.font:
            target.font = copy(source.font)
        if source.protection:
            target.protection = copy(source.protection)


def _orders_to_cio_rows(client_code: str, orders_df: pd.DataFrame) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    if orders_df is None or orders_df.empty:
        return rows

    for _, order in orders_df.iterrows():
        quantity = pd.to_numeric(_get_order_value(order, ["QTDD", "Qtd", "Qtd. Total", "quantidade"]), errors="coerce")
        if pd.isna(quantity) or int(quantity) <= 0:
            continue

        lado = str(_get_order_value(order, ["LADO", "lado"]) or "").strip().upper()
        cv = {"COMPRA": "C", "C": "C", "VENDA": "V", "V": "V"}.get(lado)
        if cv is None:
            continue

        rows.append(
            {
                "estrategia": "Simples",
                "cliente": str(client_code).strip(),
                "ativo": str(_get_order_value(order, ["ATIVO", "ticker", "Ativo"]) or "").strip(),
                "cv": cv,
                "preco": None,
                "qtdtotal": int(quantity),
            }
        )
    return rows


def generate_ordem_cio_excel(template_path, client_code, orders_df) -> bytes:
    template_path = Path(template_path)
    wb = load_workbook(template_path)
    ws = wb.active

    header_row, columns = _find_ordem_cio_headers(ws)
    data_start_row = header_row + 1
    target_columns = [
        columns["estrategia"],
        columns["cliente"],
        columns["ativo"],
        columns["cv"],
        columns["preco"],
        columns["qtdtotal"],
    ]

    rows = _orders_to_cio_rows(client_code, orders_df)
    last_row = max(ws.max_row, data_start_row + len(rows) - 1)
    style_source_row = data_start_row

    for row_idx in range(data_start_row, last_row + 1):
        if row_idx > style_source_row:
            _copy_row_style(ws, style_source_row, row_idx, target_columns)
        for col_idx in target_columns:
            ws.cell(row=row_idx, column=col_idx).value = None

    for offset, order in enumerate(rows):
        row_idx = data_start_row + offset
        if row_idx > style_source_row:
            _copy_row_style(ws, style_source_row, row_idx, target_columns)
        ws.cell(row=row_idx, column=columns["estrategia"]).value = order["estrategia"]
        ws.cell(row=row_idx, column=columns["cliente"]).value = order["cliente"]
        ws.cell(row=row_idx, column=columns["ativo"]).value = order["ativo"]
        ws.cell(row=row_idx, column=columns["cv"]).value = order["cv"]
        ws.cell(row=row_idx, column=columns["preco"]).value = None
        ws.cell(row=row_idx, column=columns["qtdtotal"]).value = order["qtdtotal"]

    output = BytesIO()
    wb.save(output)
    filled_path = template_path.parent / "ordem_cio_preenchida.xlsx"
    try:
        wb.save(filled_path)
    except OSError:
        pass
    return output.getvalue()
